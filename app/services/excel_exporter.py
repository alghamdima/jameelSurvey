"""خدمة تصدير نتائج الاستبيانات إلى ملف إكسل رسمي (.xlsx) بهوية عبد اللطيف جميل للتمويل."""
import io
import json
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Survey, SurveyResponse

def safe_cell(sheet, row, column, value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def generate_survey_excel(survey: Survey, responses: list[SurveyResponse], stats: dict, lang: str = "ar") -> io.BytesIO:
    """توليد ملف Excel احترافي كامل يحتوي على ملخص تفصيلي وسجل الإجابات الفردية."""
    wb = openpyxl.Workbook()

    # ألوان هوية عبد اللطيف جميل للتمويل
    ALJ_PURPLE = "3C1053"
    ALJ_HEADER_PURPLE = "4A1566"
    ALJ_GOLD = "FFCD00"
    ALJ_LIGHT_GRAY = "F8F9FA"
    ALJ_BORDER_GRAY = "E2E8F0"
    ALJ_TEXT_DARK = "1E293B"

    header_fill = PatternFill(start_color=ALJ_PURPLE, end_color=ALJ_PURPLE, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    sub_header_fill = PatternFill(start_color=ALJ_HEADER_PURPLE, end_color=ALJ_HEADER_PURPLE, fill_type="solid")
    sub_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    highlight_gold_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=ALJ_BORDER_GRAY),
        right=Side(style="thin", color=ALJ_BORDER_GRAY),
        top=Side(style="thin", color=ALJ_BORDER_GRAY),
        bottom=Side(style="thin", color=ALJ_BORDER_GRAY),
    )

    is_rtl = (lang == "ar")

    # ----------------------------------------------------
    # الورقة 1: الملخص الإحصائي المتقدم (Overview & Summary)
    # ----------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "ملخص التحليلات" if is_rtl else "Analytics Summary"
    ws_summary.views.sheetView[0].rightToLeft = is_rtl

    # عنوان التقرير الرئيسي
    title_text = survey.title_ar if is_rtl else survey.title_en
    ws_summary.merge_cells("A1:G1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"تقرير نتائج الاستبيان: {title_text}" if is_rtl else f"Survey Analytics Report: {title_text}"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color=ALJ_PURPLE)
    title_cell.alignment = Alignment(horizontal="right" if is_rtl else "left", vertical="center")
    ws_summary.row_dimensions[1].height = 30

    # بطاقات المؤشرات العامة (KPIs)
    kpis = [
        ("إجمالي الموظفين المشاركين" if is_rtl else "Total Participants", len(responses)),
        ("عدد الأسئلة" if is_rtl else "Questions Count", len(survey.questions)),
        ("تاريخ الاستخراج" if is_rtl else "Export Date", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("حالة الاستبيان" if is_rtl else "Survey Status", survey.status.upper()),
    ]
    
    for idx, (lbl, val) in enumerate(kpis, start=1):
        c_lbl = ws_summary.cell(row=3, column=idx)
        c_lbl.value = lbl
        c_lbl.font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
        c_lbl.fill = zebra_fill
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_lbl.border = thin_border

        c_val = ws_summary.cell(row=4, column=idx)
        c_val.value = val
        c_val.font = Font(name="Segoe UI", size=12, bold=True, color=ALJ_PURPLE)
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        c_val.border = thin_border

    ws_summary.row_dimensions[3].height = 20
    ws_summary.row_dimensions[4].height = 26

    # جدول تفصيل إحصائيات الأسئلة والخيارات
    headers = [
        ("#", 6),
        ("نص السؤال بالعربية" if is_rtl else "Question (Arabic)", 35),
        ("Question (English)", 35),
        ("نوع السؤال" if is_rtl else "Question Type", 18),
        ("الخيار / الإجابة" if is_rtl else "Choice / Answer", 32),
        ("عدد الأصوات" if is_rtl else "Vote Count", 14),
        ("النسبة المئوية" if is_rtl else "Percentage", 14),
    ]

    start_row = 6
    ws_summary.row_dimensions[start_row].height = 24
    for col_idx, (h_name, width) in enumerate(headers, start=1):
        c = ws_summary.cell(row=start_row, column=col_idx)
        c.value = h_name
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws_summary.column_dimensions[col_letter].width = width

    current_r = start_row + 1
    total_resp_count = len(responses)

    for q_idx, q in enumerate(survey.questions, start=1):
        item = stats.get(q.question_key, {})
        opt_counts = item.get("opt_counts", {})
        top_opt_key = item.get("top_opt_key")
        q_type_name = "اختيار فردي" if q.question_type == "single_choice" else ("اختيار متعدد" if q.question_type == "multiple_choice" else "نصي حر")
        if not is_rtl:
            q_type_name = "Single Choice" if q.question_type == "single_choice" else ("Multiple Choice" if q.question_type == "multiple_choice" else "Free Text")

        if q.question_type in ["single_choice", "multiple_choice"]:
            opts = q.get_options()
            for opt_idx, opt in enumerate(opts):
                cnt = opt_counts.get(opt["key"], 0)
                pct_val = round((cnt / total_resp_count * 100), 1) if total_resp_count > 0 else 0
                is_top = (top_opt_key == opt["key"] and cnt > 0)

                row_vals = [
                    q_idx if opt_idx == 0 else "",
                    q.text_ar if opt_idx == 0 else "",
                    q.text_en if opt_idx == 0 else "",
                    q_type_name if opt_idx == 0 else "",
                    f"⭐ {opt.get('text_ar' if is_rtl else 'text_en', '')}" if is_top else opt.get('text_ar' if is_rtl else 'text_en', ''),
                    cnt,
                    f"{pct_val}%",
                ]

                ws_summary.row_dimensions[current_r].height = 22
                for c_i, v in enumerate(row_vals, start=1):
                    cell = safe_cell(ws_summary, current_r, c_i, v)
                    cell.border = thin_border
                    cell.font = Font(name="Segoe UI", size=10, bold=(is_top or c_i == 1))
                    if is_top:
                        cell.fill = highlight_gold_fill
                    elif current_r % 2 == 0:
                        cell.fill = zebra_fill
                    
                    if c_i in [1, 4, 6, 7]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right" if is_rtl else "left", vertical="center")

                current_r += 1
        else:
            # سؤال نصي
            text_answers = item.get("text_answers", [])
            answered_cnt = len(text_answers)
            pct_val = round((answered_cnt / total_resp_count * 100), 1) if total_resp_count > 0 else 0
            
            row_vals = [
                q_idx,
                q.text_ar,
                q.text_en,
                q_type_name,
                f"إجمالي الإجابات النصية: {answered_cnt}" if is_rtl else f"Total Text Answers: {answered_cnt}",
                answered_cnt,
                f"{pct_val}%",
            ]
            ws_summary.row_dimensions[current_r].height = 22
            for c_i, v in enumerate(row_vals, start=1):
                cell = safe_cell(ws_summary, current_r, c_i, v)
                cell.border = thin_border
                cell.font = Font(name="Segoe UI", size=10, bold=(c_i == 1))
                if current_r % 2 == 0:
                    cell.fill = zebra_fill
                if c_i in [1, 4, 6, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right" if is_rtl else "left", vertical="center")
            current_r += 1

    # ----------------------------------------------------
    # الورقة 2: سجل الإجابات التفصيلي للموظفين (Detailed Raw Responses)
    # ----------------------------------------------------
    ws_raw = wb.create_sheet(title="سجل المشاركات" if is_rtl else "Raw Responses")
    ws_raw.views.sheetView[0].rightToLeft = is_rtl

    # ترويسة سجل المشاركات
    raw_headers = [
        ("#", 6),
        ("الرقم الوظيفي" if is_rtl else "Employee ID", 18),
        ("تاريخ وساعة الإرسال" if is_rtl else "Submitted At", 22),
    ]

    # إضافة عمود لكل سؤال
    for q_i, q in enumerate(survey.questions, start=1):
        q_label = f"س{q_i}: {q.text_ar[:30]}..." if is_rtl else f"Q{q_i}: {q.text_en[:30]}..."
        raw_headers.append((q_label, 30))

    ws_raw.row_dimensions[1].height = 25
    for c_idx, (h_title, w) in enumerate(raw_headers, start=1):
        c = ws_raw.cell(row=1, column=c_idx, value=h_title)
        c.fill = sub_header_fill
        c.font = sub_header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws_raw.column_dimensions[col_letter].width = w

    # إدراج صفوف المشاركين
    for r_idx, resp in enumerate(responses, start=1):
        row_num = r_idx + 1
        ws_raw.row_dimensions[row_num].height = 20
        ans_data = {}
        try:
            ans_data = json.loads(resp.answers_json) if resp.answers_json else {}
        except Exception:
            pass

        row_content = [
            r_idx,
            resp.employee_id,
            resp.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if resp.submitted_at else "",
        ]

        # فك شفرة كل إجابة وربطها بنصوص الخيارات المفهومة
        for q in survey.questions:
            user_ans = ans_data.get(q.question_key)
            if not user_ans:
                row_content.append("-" if is_rtl else "-")
            elif q.question_type == "single_choice":
                # العثور على نص الخيار
                opts = q.get_options()
                opt_match = next((o for o in opts if o["key"] == user_ans), None)
                val_txt = opt_match.get("text_ar" if is_rtl else "text_en", user_ans) if opt_match else str(user_ans)
                row_content.append(val_txt)
            elif q.question_type == "multiple_choice":
                opts = q.get_options()
                if isinstance(user_ans, list):
                    labels = []
                    for k in user_ans:
                        m = next((o for o in opts if o["key"] == k), None)
                        labels.append(m.get("text_ar" if is_rtl else "text_en", k) if m else str(k))
                    row_content.append(", ".join(labels))
                else:
                    row_content.append(str(user_ans))
            else:
                row_content.append(str(user_ans).strip())

        for c_i, v in enumerate(row_content, start=1):
            cell = safe_cell(ws_raw, row_num, c_i, v)
            cell.border = thin_border
            cell.font = Font(name="Segoe UI", size=10)
            if row_num % 2 == 0:
                cell.fill = zebra_fill
            if c_i in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right" if is_rtl else "left", vertical="center")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
